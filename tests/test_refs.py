"""--ref: formulas referencing given ranges."""

import json
import os
from pathlib import Path

import openpyxl
import pytest
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from xlgrep.address import CellRange, parse_range
from xlgrep.cli import main
from xlgrep.refs import RefFinder, Target, parse_target, scan_refs


def refs_of(formula):
    return [(formula[r.start : r.end], r.sheets, r.external, r.name) for r in scan_refs(formula)]


def test_scan_refs_forms():
    assert refs_of("=SUM(Data!A1:A10)+'Raw Data'!$B$2") == [
        ("Data!A1:A10", ("Data",), False, None),
        ("'Raw Data'!$B$2", ("Raw Data",), False, None),
    ]
    assert refs_of("=VLOOKUP(A2,Data!A:D,4,0)+SUM(3:3)") == [
        ("A2", None, False, None), ("Data!A:D", ("Data",), False, None), ("3:3", None, False, None),
    ]
    assert refs_of("=SUM(Sheet1:Sheet3!B2)") == [("Sheet1:Sheet3!B2", ("Sheet1", "Sheet3"), False, None)]
    assert refs_of("=[1]Ext!A1+'[Book.xlsx]Ext'!B2") == [
        ("[1]Ext!A1", ("Ext",), True, None), ("'[Book.xlsx]Ext'!B2", ("Ext",), True, None),
    ]
    assert refs_of("=Data!A1#+LOG10(5)") == [("Data!A1#", ("Data",), False, None)]


def test_scan_refs_ignores_strings_tables_and_function_names():
    assert refs_of('=INDIRECT("Data!A1")') == []
    assert refs_of("=SUM(Table1[Amount])+Table1[[#This Row],[Qty]]") == []
    assert refs_of("=_xlfn.XLOOKUP(A1,B:B,C:C)") == [("A1", None, False, None), ("B:B", None, False, None),
                                                    ("C:C", None, False, None)]
    assert refs_of("=TaxRate*A1:INDEX(B:B,5)+TRUE") == [
        ("TaxRate", None, False, "TaxRate"), ("A1", None, False, None), ("B:B", None, False, None),
    ]


def test_parse_target():
    assert parse_target("Data!A:D") == Target("Data", CellRange(None, 1, None, 4))
    assert parse_target("'Raw Data'!B2:F10") == Target("Raw Data", CellRange(2, 2, 10, 6))
    assert parse_target("Data!") == Target("Data", None)
    assert parse_target("A1:B5") == Target(None, CellRange(1, 1, 5, 2))
    with pytest.raises(ValueError):
        parse_target("Data!nope")


def test_ref_finder_names_3d_and_sweep():
    finder = RefFinder(
        [parse_target("Data!A:D")],
        ["Summary", "Data", "Other"],
        {(None, "SALESRANGE"): "=Data!B:B", (None, "CHAINED"): "=SalesRange*2", (None, "LOOP"): "=Loop+1",
         ("Summary", "LOCAL"): "=Other!A1"},
    )
    assert [(h.text, h.via) for h in finder.find("=SUM(SalesRange)", "Summary")] == [("SalesRange", "SalesRange")]
    assert [h.via for h in finder.find("=Chained", "Summary")] == ["Chained"]
    assert finder.find("=Loop", "Summary") == []
    assert finder.find("=Local", "Summary") == []
    assert [h.text for h in finder.find("=A2*2", "Data")] == ["A2"]
    assert finder.find("=A2*2", "Summary") == []
    assert [h.text for h in finder.find("=SUM(Summary:Other!C5)", "Summary")] == ["Summary:Other!C5"]
    assert finder.find("=Data!E1", "Summary") == []

    sweep = RefFinder([parse_target("Summary!A40")], ["Summary"], {})
    assert [h.text for h in sweep.find("=$A2>1", "Summary", [parse_range("B2:B50")])] == ["$A2"]
    assert sweep.find("=$A$2>1", "Summary", [parse_range("B2:B50")]) == []


@pytest.fixture
def run(tmp_path: Path, capsys):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "=VLOOKUP(A2,Data!A:D,4,FALSE)"
    ws["A2"] = "=SUM(Other!A1:A9)"
    ws["A3"] = "=SUM(SalesRange)"
    ws["A4"] = '=INDIRECT("Data!A1")'
    ws["A5"] = "=[1]Data!A1"
    ws["A6"] = "=SUM(Summary:Other!B2)"
    ws["A7"] = "Data!A1 as text"
    ws.conditional_formatting.add("C2:C50", FormulaRule(formula=["$B2>Data!$E$1"]))
    dv = DataValidation(type="list", formula1="=Data!$C$1:$C$9")
    dv.add("D1:D10")
    ws.add_data_validation(dv)
    wb.defined_names["SalesRange"] = DefinedName("SalesRange", attr_text="Data!$B:$B")
    data = wb.create_sheet("Data")
    data["F1"] = "=A1*2"
    data["F2"] = "=SUM(H1:H5)"
    other = wb.create_sheet("Other")
    other["A1"] = "=Summary!A1"
    wb.save(tmp_path / "refs.xlsx")

    def _run(*args):
        cwd = os.getcwd()
        os.chdir(tmp_path)
        try:
            code = main(["--color", "never", *args])
        finally:
            os.chdir(cwd)
        out, err = capsys.readouterr()
        return code, out, err

    return _run


def test_ref_columns(run):
    code, out, _ = run("--ref", "Data!A:D")
    assert code == 0
    assert out.splitlines() == [
        "refs.xlsx:Summary!A1:=VLOOKUP(A2,Data!A:D,4,FALSE)",
        "refs.xlsx:Summary!A3:=SUM(SalesRange)",
        "refs.xlsx:Summary!A6:=SUM(Summary:Other!B2)",
        "refs.xlsx:Summary!D1:D10#dv:=Data!$C$1:$C$9",
        "refs.xlsx:Data!F1:=A1*2",
        "refs.xlsx:SalesRange#name:=Data!$B:$B",
    ]


def test_ref_whole_sheet_and_cf_sweep(run):
    _, out, _ = run("--ref", "Data!", "--objects", "cf")
    assert out == "refs.xlsx:Summary!C2:C50#cf:=$B2>Data!$E$1\n"
    _, out, _ = run("--ref", "Summary!B30", "--objects", "cf")
    assert out == "refs.xlsx:Summary!C2:C50#cf:=$B2>Data!$E$1\n"
    _, out, _ = run("--ref", "Summary!B51", "--objects", "cf")
    assert out == ""


def test_ref_any_sheet_and_multiple(run):
    _, out, _ = run("--ref", "H3", "--ref", "Other!A5", "--objects", "cells")
    assert out.splitlines() == [
        "refs.xlsx:Summary!A2:=SUM(Other!A1:A9)",
        "refs.xlsx:Data!F2:=SUM(H1:H5)",
    ]


def test_ref_and_pattern(run):
    _, out, _ = run("-f", "SUM", "--ref", "Data!A:D")
    assert out.splitlines() == [
        "refs.xlsx:Summary!A3:=SUM(SalesRange)",
        "refs.xlsx:Summary!A6:=SUM(Summary:Other!B2)",
    ]
    _, out, _ = run("-c", "--ref", "Data!A:D", "-e", "VLOOKUP")
    assert out == "refs.xlsx:1\n"


def test_ref_json_and_highlight(run):
    _, out, _ = run("--json", "--ref", "Data!B:B", "--objects", "cells", "--range", "A3")
    rec = json.loads(out)
    assert rec["cell"] == "A3"
    assert rec["refs"] == [{"start": 5, "end": 15, "text": "SalesRange", "via": "SalesRange"}]
    assert rec["matches"] == [{"start": 5, "end": 15, "text": "SalesRange"}]
    _, out, _ = run("--color", "always", "--ref", "Other!A1", "--objects", "cells", "--sheet", "Summary")
    assert "SUM(\x1b[1;31mOther!A1:A9\x1b[0m)" in out


def test_ref_errors(run):
    with pytest.raises(SystemExit):
        run("--ref", "Data!nope")
    with pytest.raises(SystemExit):
        run("--ref", "Data!A1", "--in", "value")
    with pytest.raises(SystemExit):
        run("--list-funcs", "--ref", "Data!A1")
