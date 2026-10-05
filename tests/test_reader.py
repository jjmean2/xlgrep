"""The direct cell reader: regex path, XML-parser fallback and openpyxl must agree."""

import datetime as dt
import warnings
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from xlgrep.package import open_package
from xlgrep.workbook import _context, read_cells, read_sheets

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

SHARED_STRINGS = [
    '<si><r><t>Hel</t></r><r><rPr><b/></rPr><t>lo</t></r><rPh sb="0" eb="1"><t>ignored</t></rPh></si>',
    "<si><t>line_x000D_break</t></si>",
    "<si><t>한글 &amp; 기호</t></si>",
]

STYLES = (
    f'<styleSheet xmlns="{MAIN}"><numFmts count="1"><numFmt numFmtId="164" formatCode="yyyy-mm-dd hh:mm"/></numFmts>'
    '<cellXfs count="5"><xf numFmtId="0"/><xf numFmtId="14"/><xf numFmtId="20"/><xf numFmtId="164"/>'
    '<xf numFmtId="31"/></cellXfs></styleSheet>'
)

CELLS = [
    (
        '<row r="1">'
        '<c r="A1" t="s"><v>0</v></c><c r="B1"><v>42</v></c><c r="C1" t="b"><v>1</v></c>'
        '<c r="D1" s="1"><v>45000</v></c><c r="E1" t="inlineStr"><is><t>a &amp; b</t></is></c>'
        '<c r="F1" t="str"><f>A1&amp;"x"</f><v>Hellox</v></c>'
        '<c r="G1"><f t="array" ref="G1:G2">SUM(B1:B2*2)</f><v>87</v></c>'
        '<c r="H1"><f>_xlfn.XLOOKUP(A1,B:B,C:C)</f><v>1</v></c><c r="J1" s="1"/>'
        '<c r="K1"><f t="dataTable" ref="K1:K2" dt2D="0" dtr="0" r1="B1"/><v>5</v></c>'
        "</row>"
    ),
    (
        '<row r="2">'
        '<c r="A2" t="s"><v>1</v></c><c r="B2"><v>1.5</v></c><c r="C2" t="e"><v>#N/A</v></c>'
        '<c r="D2" s="2"><v>0.5</v></c><c r="E2" t="inlineStr"><is><r><t>rich </t></r><r><t>inline</t></r></is></c>'
        '<c r="F2"><f t="shared" ref="F2:F4" si="0">B2*2</f><v>3</v></c><c r="G2"><v>0</v></c>'
        '<c r="H2" s="3"><v>45000.25</v></c>'
        "</row>"
    ),
    (
        '<row r="3">'
        '<c r="A3" t="s"><v>2</v></c><c r="B3"><v>1E3</v></c><c r="D3" s="4"><v>45000</v></c>'
        '<c r="F3"><f t="shared" si="0"/><v>2000</v></c>'
        "</row>"
    ),
    '<row r="4"><c r="F4"><f t="shared" si="0"/></c></row>',
]


def sheet_xml(rows=CELLS, prefix=""):
    p = f"{prefix}:" if prefix else ""
    ns = f'xmlns:{prefix}="{MAIN}"' if prefix else f'xmlns="{MAIN}"'
    body = "".join(rows)
    if prefix:
        body = body.replace("<", f"<{p}").replace(f"<{p}/", f"</{p}")
    return (
        f'<?xml version="1.0" encoding="UTF-8"?><{p}worksheet {ns}><{p}sheetData>{body}</{p}sheetData></{p}worksheet>'
    )


def write_xlsx(path: Path, sheet: str, date1904: bool = False) -> Path:
    strings = f'<sst xmlns="{MAIN}" count="{len(SHARED_STRINGS)}">{"".join(SHARED_STRINGS)}</sst>'
    workbook_pr = '<workbookPr date1904="1"/>' if date1904 else ""
    parts = {
        "[Content_Types].xml": (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '<Override PartName="/xl/sharedStrings.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
            '<Override PartName="/xl/styles.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
            "</Types>"
        ),
        "_rels/.rels": (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>'
        ),
        "xl/workbook.xml": (
            f'<workbook xmlns="{MAIN}" xmlns:r="{REL}">{workbook_pr}'
            '<sheets><sheet name="Data" sheetId="1" r:id="rId1"/></sheets></workbook>'
        ),
        "xl/_rels/workbook.xml.rels": (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{REL}/worksheet" Target="worksheets/sheet1.xml"/>'
            f'<Relationship Id="rId2" Type="{REL}/sharedStrings" Target="sharedStrings.xml"/>'
            f'<Relationship Id="rId3" Type="{REL}/styles" Target="styles.xml"/></Relationships>'
        ),
        "xl/worksheets/sheet1.xml": sheet,
        "xl/sharedStrings.xml": strings,
        "xl/styles.xml": STYLES,
    }
    with zipfile.ZipFile(path, "w") as zf:
        for name, data in parts.items():
            zf.writestr(name, data)
    return path


def cells_of(path: Path, raw_formula=True, force_etree=False):
    with zipfile.ZipFile(path) as zf:
        pkg = open_package(zf)
        ctx = _context(zf, pkg, raw_formula)
        cells = read_cells(zf.read(pkg.sheets[0].part), ctx, force_etree=force_etree)
    return {(c.row, c.col): (c.formula, c.value, c.has_cached_value) for c in cells.values()}


@pytest.fixture
def book(tmp_path):
    return write_xlsx(tmp_path / "book.xlsx", sheet_xml())


def test_values_and_formulas(book):
    cells = cells_of(book)
    assert cells[(1, 1)] == (None, "Hello", True)  # rich text runs, phonetic guide dropped
    assert cells[(2, 1)] == (None, "line\rbreak", True)  # _x000D_
    assert cells[(3, 1)] == (None, "한글 & 기호", True)
    assert cells[(1, 2)] == (None, 42, True)
    assert cells[(2, 2)] == (None, 1.5, True)
    assert cells[(3, 2)] == (None, 1000.0, True)
    assert cells[(1, 3)] == (None, True, True)
    assert cells[(2, 3)] == (None, "#N/A", True)
    assert cells[(1, 4)] == (None, dt.datetime(2023, 3, 15), True)
    assert cells[(2, 4)] == (None, dt.time(12, 0), True)
    assert cells[(3, 4)] == (None, dt.datetime(2023, 3, 15), True)  # CJK built-in date id 31
    assert cells[(2, 8)] == (None, dt.datetime(2023, 3, 15, 6, 0), True)
    assert cells[(1, 5)] == (None, "a & b", True)
    assert cells[(2, 5)] == (None, "rich inline", True)
    assert cells[(1, 6)] == ('=A1&"x"', "Hellox", True)
    assert cells[(2, 6)] == ("=B2*2", 3, True)
    assert cells[(3, 6)] == ("=B3*2", 2000, True)  # shared formula expanded
    assert cells[(4, 6)] == ("=B4*2", None, False)  # no cached value
    assert cells[(1, 7)] == ("=SUM(B1:B2*2)", 87, True)
    assert cells[(2, 7)] == (None, 0, True)
    assert cells[(1, 8)] == ("=_xlfn.XLOOKUP(A1,B:B,C:C)", 1, True)
    assert cells[(1, 11)] == ("=TABLE(B1)", 5, True)
    assert (1, 10) not in cells  # styled but empty


def test_normalized_formulas(book):
    assert cells_of(book, raw_formula=False)[(1, 8)][0] == "=XLOOKUP(A1,B:B,C:C)"


def test_regex_and_parser_paths_agree(book):
    assert cells_of(book) == cells_of(book, force_etree=True)


def test_prefixed_namespace_falls_back(tmp_path, book):
    prefixed = write_xlsx(tmp_path / "prefixed.xlsx", sheet_xml(prefix="x"))
    assert cells_of(prefixed) == cells_of(book)


def test_implicit_coordinates_fall_back(tmp_path, book):
    rows = [
        '<row><c t="s"><v>0</v></c><c><v>42</v></c></row>',
        '<row r="3"><c r="B3"><v>7</v></c><c t="b"><v>0</v></c></row>',
    ]
    cells = cells_of(write_xlsx(tmp_path / "implicit.xlsx", sheet_xml(rows)))
    assert cells == {
        (1, 1): (None, "Hello", True),
        (1, 2): (None, 42, True),
        (3, 2): (None, 7, True),
        (3, 3): (None, False, True),
    }


def test_date1904(tmp_path):
    rows = ['<row r="1"><c r="A1" s="1"><v>0</v></c><c r="B1" s="1"><v>1</v></c></row>']
    cells = cells_of(write_xlsx(tmp_path / "mac.xlsx", sheet_xml(rows), date1904=True))
    assert cells[(1, 2)][1] == dt.datetime(1904, 1, 2)


def test_matches_openpyxl(book):
    """Everything openpyxl reads the same way (it keeps _xHHHH_ escapes and skips CJK date ids)."""
    ours = cells_of(book)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        formulas = openpyxl.load_workbook(book, read_only=True)["Data"]
        values = openpyxl.load_workbook(book, read_only=True, data_only=True)["Data"]
        f_cells = {
            (c.row, c.column): c.value
            for row in formulas.iter_rows()
            for c in row
            if c.value is not None and hasattr(c, "row")
        }
        v_cells = {
            (c.row, c.column): c.value
            for row in values.iter_rows()
            for c in row
            if c.value is not None and hasattr(c, "row")
        }
    skip = {(2, 1), (3, 4)}
    for pos, raw in f_cells.items():
        if pos in skip:
            continue
        formula, value, _ = ours[pos]
        if isinstance(raw, ArrayFormula):
            assert formula == raw.text
        elif isinstance(raw, DataTableFormula):
            assert formula.startswith("=TABLE(")
        elif isinstance(raw, str) and raw.startswith("="):
            assert formula == raw, pos
        else:
            assert (formula, value) == (None, raw), pos
        if pos in v_cells:
            assert value == v_cells[pos], pos
    assert set(ours) - skip == set(f_cells) - skip


def test_read_sheets_on_openpyxl_files(sample_dir):
    """Workbooks written by openpyxl (inline strings, no shared strings) read fine."""
    sheets = {s.name: s for s in read_sheets(sample_dir / "sales.xlsx")}
    assert sheets["Summary"].get(2, 2).formula == "=VLOOKUP(A2,Data!A:D,4,FALSE)"
    assert sheets["Summary"].get(2, 2).value == 1200
    assert sheets["Raw Data"].get(2, 1).value == "line one\nline two"


def test_formulas_only(book):
    with zipfile.ZipFile(book) as zf:
        pkg = open_package(zf)
        ctx = _context(zf, pkg, True)
        data = zf.read(pkg.sheets[0].part)
        for force in (False, True):
            cells = read_cells(data, ctx, force_etree=force, formulas_only=True)
            assert sorted(cells) == [(1, 6), (1, 7), (1, 8), (1, 11), (2, 6), (3, 6), (4, 6)]
