"""--deps: dependencies on other workbooks, between sheets, and on data connections."""

import json
import os
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName

from xlgrep.cli import main
from xlgrep.deps import MATCHED, MISSING, resolve_target

REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def add_links(path: Path, targets: list[str], connection: str | None = None) -> None:
    """Make [1], [2], … in formulas point at ``targets``, like Excel stores external links."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename).decode()
            if item.filename == "xl/workbook.xml":
                refs = "".join(f'<externalReference r:id="rIdX{i}"/>' for i in range(1, len(targets) + 1))
                anchor = "<definedNames>" if "<definedNames>" in data else "<calcPr"
                data = data.replace(anchor, f"<externalReferences>{refs}</externalReferences>{anchor}", 1)
            elif item.filename == "xl/_rels/workbook.xml.rels":
                extra = "".join(
                    f'<Relationship Id="rIdX{i}" Type="{REL}/externalLink" Target="externalLinks/externalLink{i}.xml"/>'
                    for i in range(1, len(targets) + 1)
                )
                if connection:
                    extra += f'<Relationship Id="rIdConn" Type="{REL}/connections" Target="connections.xml"/>'
                data = data.replace("</Relationships>", extra + "</Relationships>")
            dst.writestr(item, data)
        for i, target in enumerate(targets, 1):
            dst.writestr(f"xl/externalLinks/externalLink{i}.xml",
                         f'<externalLink xmlns="{MAIN}" xmlns:r="{REL}"><externalBook r:id="rId1">'
                         '<sheetNames><sheetName val="Budget"/></sheetNames></externalBook></externalLink>')
            dst.writestr(f"xl/externalLinks/_rels/externalLink{i}.xml.rels",
                         '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                         f'<Relationship Id="rId1" Type="{REL}/externalLinkPath" Target="{target}" '
                         'TargetMode="External"/></Relationships>')
        if connection:
            dst.writestr("xl/connections.xml",
                         f'<connections xmlns="{MAIN}"><connection id="1" name="{connection}" type="5">'
                         '<dbPr connection="Provider=Microsoft.Mashup.OleDb.1;Password=secret" command=""/>'
                         "</connection></connections>")
    shutil.move(tmp, path)


@pytest.fixture
def tree(tmp_path: Path) -> Path:
    (tmp_path / "sub").mkdir()
    for name in ("budget.xlsx", "sub/lookup.xlsx"):
        wb = openpyxl.Workbook()
        wb.active["A1"] = 1
        wb.save(tmp_path / name)

    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "Summary"
    data = wb.create_sheet("Data")
    report = wb.create_sheet("Report")
    summary["A1"] = "=[1]Budget!B2*2"
    summary["A2"] = "=SUM([1]Budget!A:A)+Data!A1"
    summary["A3"] = "=VLOOKUP(A1,Rates,2,0)"
    summary["A4"] = "=Data!B1+Data!C1"  # one formula, one dependency on Data
    summary["A5"] = "=SUM(Summary:Report!Z1)"  # 3D: Data and Report (Summary itself isn't a dependency)
    report["A1"] = "=Summary!A1"
    report["A2"] = "=[3]Sheet1!A1"
    data["A1"] = 5
    data.conditional_formatting.add("A1:A9", FormulaRule(formula=["Summary!$A$1>0"]))
    wb.defined_names["Rates"] = DefinedName("Rates", attr_text="[2]Rates!$A:$B")
    wb.save(tmp_path / "sales.xlsx")
    add_links(tmp_path / "sales.xlsx",
              ["budget.xlsx", "file:///C:\\%EA%B3%B5%EC%9C%A0\\rates.xlsx", "file:///D:\\old\\Lookup.xlsx"],
              connection="Query - Sales")
    return tmp_path


@pytest.fixture
def run(tree, capsys):
    def _run(*args):
        cwd = os.getcwd()
        os.chdir(tree)
        try:
            code = main(["--color", "never", "--deps", *args])
        finally:
            os.chdir(cwd)
        out, err = capsys.readouterr()
        return code, out, err

    return _run


def test_text(run):
    code, out, _ = run()
    assert code == 0
    assert out.splitlines() == [
        "sales.xlsx",
        "  → C:\\공유\\rates.xlsx      2 formulas  (missing)",
        "  → budget.xlsx             2 formulas",
        "  → sub/lookup.xlsx          1 formula  (matched by name)",
        "  sheets  Summary → Data 3 · Summary → Report 1 · Data → Summary 1 · Report → Summary 1",
        "  data    Query - Sales (Power Query)",
        "1 of 3 workbooks have dependencies",
    ]


def test_by_sheet(run):
    _, out, _ = run("--by", "sheet")
    assert out.splitlines() == [
        "sales.xlsx",
        "  (workbook) → C:\\공유\\rates.xlsx : Rates       1 formula  (missing)",
        "  Summary → C:\\공유\\rates.xlsx : Rates          1 formula  (missing)",
        "  Summary → budget.xlsx : Budget               2 formulas",
        "  Report → sub/lookup.xlsx : Sheet1             1 formula  (matched by name)",
        "  Summary → Data                               3 formulas",
        "  Summary → Report                              1 formula",
        "  Data → Summary                                1 formula",
        "  Report → Summary                              1 formula",
        "  data    Query - Sales (Power Query)",
        "1 of 3 workbooks have dependencies",
    ]


def test_json_csv_and_no_secrets(run):
    _, out, _ = run("--json")
    records = [json.loads(line) for line in out.splitlines()]
    assert {"from": "sales.xlsx", "from_sheet": None, "to": "budget.xlsx", "to_sheet": None, "kind": "file",
            "formulas": 2, "status": "found", "type": None} in records
    assert {"from": "sales.xlsx", "from_sheet": None, "to": "Query - Sales", "to_sheet": None, "kind": "data",
            "formulas": None, "status": None, "type": "Power Query"} in records
    assert "secret" not in out
    _, out, _ = run("--csv")
    assert out.splitlines()[0] == "from,from_sheet,to,to_sheet,kind,formulas,status,type"


def test_mermaid(run):
    _, out, _ = run("--graph", "mermaid")
    lines = out.splitlines()
    assert lines[0] == "flowchart LR"
    assert '  n1["rates.xlsx (missing)"]:::missing' in lines
    assert "  n0 -.->|2| n1" in lines
    assert "  n0 -->|2| n2" in lines
    assert '  n4[("Query - Sales (Power Query)")]' in lines
    _, out, _ = run("--graph", "mermaid", "--by", "sheet")
    lines = out.splitlines()
    box = lines.index('  subgraph w0["sales.xlsx"]')
    assert lines[box + 1 : box + 6] == ['    n0["(workbook)"]', '    n2["Summary"]', '    n4["Report"]',
                                        '    n6["Data"]', "  end"]
    assert "  n2 -->|3| n6" in lines and "  w0 --> n7" in lines


def test_parallel_and_exit_codes(run):
    assert run("-j", "2") == run("-j", "1")
    code, out, _ = run("budget.xlsx")
    assert (code, out) == (1, "")
    for args in (["-f", "SUM"], ["--graph", "mermaid", "--stats"]):
        with pytest.raises(SystemExit):
            run(*args)


def test_resolve_target(tmp_path):
    (tmp_path / "a.xlsx").write_bytes(b"")
    searched = [tmp_path / "a.xlsx", tmp_path / "b.xlsx"]
    assert resolve_target("file:///C:\\x\\A.XLSX", tmp_path / "s.xlsx", searched)[1] == MATCHED
    assert resolve_target("file://server/share/z.xlsx", tmp_path / "s.xlsx", searched) == (
        "\\\\server/share/z.xlsx", MISSING)
