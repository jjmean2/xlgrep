"""--stats: workbook metrics."""

import json
import os
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table
from test_reader import sheet_xml, write_xlsx

from xlgrep.cli import main
from xlgrep.scope import Scope
from xlgrep.stats import StatsConfig, combine, stats_file

ALL = {"cell", "name", "cf", "dv", "note"}

CELL_ROWS = [
    '<row r="1">'
    '<c r="A1"><f t="shared" ref="A1:A3" si="0">B1*2</f><v>2</v></c>'
    '<c r="D1"><f>C1*2</f><v>0</v></c>'
    '<c r="F1"><f>INDIRECT("A1")+NOW()</f><v>1</v></c>'
    '<c r="G1"><f t="array" ref="G1:G2">SUM(B1:B2*2)</f><v>4</v></c>'
    '<c r="H1"><f t="dataTable" ref="H1:H2" dt2D="0" dtr="0" r1="B1"/><v>5</v></c>'
    '<c r="J1" t="e"><v>#REF!</v></c>'
    "</row>",
    '<row r="2"><c r="A2"><f t="shared" si="0"/><v>4</v></c><c r="B2"><v>2</v></c>'
    '<c r="F2"><f>_xlfn.RANDARRAY(3)</f><v>0.5</v></c>'
    '<c r="J2" t="e"><f>1/0</f><v>#DIV/0!</v></c></row>',
    '<row r="3"><c r="A3"><f t="shared" si="0"/><v>6</v></c><c r="B3" t="s"><v>0</v></c></row>',
    '<row r="5"><c r="E5"><f>D5*2</f><v>0</v></c></row>',
]


def test_cell_metrics(tmp_path):
    path = write_xlsx(tmp_path / "cells.xlsx", sheet_xml(CELL_ROWS).replace(
        "<sheetData>", '<dimension ref="A1:J5"/><sheetData>'))
    s = stats_file(path, StatsConfig(Scope(ALL), None)).sheets[0]
    assert s.dimension == "A1:J5"
    assert (s.values, s.formulas) == (3, 10)
    # A1:A3 share one formula; D1 (=C1*2) and E5 (=D5*2) are the same logic: =RC[-1]*2.
    assert len(s.unique) == 7
    assert "=RC[-1]*2" in s.unique and "=RC[1]*2" in s.unique
    assert s.volatile_formulas == 2
    assert s.volatile == {"INDIRECT": 1, "NOW": 1, "RANDARRAY": 1}
    assert (s.arrays, s.data_tables) == (1, 1)
    assert s.errors == {"#REF!": 1, "#DIV/0!": 1}
    assert s.longest == len('=INDIRECT("A1")+NOW()')


@pytest.fixture
def book(tmp_path) -> Path:
    path = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Main"
    for row in (["key", "val"], ["k1", 1], ["k2", 2]):
        ws.append(row)
    ws["C1"] = "=MyLambda(2)+GetRate(B1)"
    ws["C2"] = "=SUM(B1:B3)"
    ws["D1"].comment = Comment("check", "me")
    ws.add_table(Table(displayName="T1", ref="A1:B3"))
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3))
    ws.add_chart(chart, "F2")
    ws.conditional_formatting.add("B1:B3", CellIsRule(operator="between", formula=["1", "SUM($B$1:$B$3)"]))
    ws.conditional_formatting.add("A1:A3", FormulaRule(formula=["LEN(A1)>2"]))
    dv = DataValidation(type="whole", formula1="0")
    dv.add("B1:B3")
    ws.add_data_validation(dv)
    wb.defined_names["MyLambda"] = DefinedName("MyLambda", attr_text="_xlfn.LAMBDA(_xlpm.x,_xlpm.x*2)")
    ws.defined_names["Local"] = DefinedName("Local", attr_text="Main!$A$1")
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "#N/A"  # a text cell that reads like an error: counted as one (documented heuristic)
    wb.save(path)

    pivot_rel = ('<Relationship Id="rIdPivot" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                 'relationships/pivotTable" Target="../pivotTables/pivotTable1.xml"/>')
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/_rels/sheet1.xml.rels":
                data = data.replace(b"</Relationships>", pivot_rel.encode() + b"</Relationships>")
            dst.writestr(item, data)
        dst.writestr("xl/pivotTables/pivotTable1.xml", "<pivotTableDefinition/>")
        dst.writestr("xl/vbaProject.bin", b"vba")
        dst.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")
        dst.writestr("xl/connections.xml", "<connections/>")
    shutil.move(tmp, path)
    return path


def test_package_metrics(book):
    f = stats_file(book, StatsConfig(Scope(ALL), None))
    assert [(s.name, s.hidden) for s in f.sheets] == [("Main", False), ("Hidden", True)]
    main_sheet = f.sheets[0]
    assert (main_sheet.cf, main_sheet.dv, main_sheet.names, main_sheet.notes) == (2, 1, 1, 1)
    assert (main_sheet.tables, main_sheet.pivots, main_sheet.charts) == (1, 1, 1)
    assert main_sheet.custom == {"GetRate"} and main_sheet.lambdas == {"MyLambda"}
    assert (f.vba, f.external_links, f.connections, f.hidden_sheets, f.workbook_names) == (True, 1, True, 1, 1)
    total = f.total
    assert total.names == 2
    assert total.errors == {"#N/A": 1}


def test_combine_unions_unique_formulas():
    from xlgrep.stats import Stats

    a, b = Stats("a", formulas=2, unique={"=RC[-1]"}), Stats("b", formulas=3, unique={"=RC[-1]", "=R1C1"})
    total = combine("t", [a, b])
    assert (total.formulas, len(total.unique)) == (5, 2)


@pytest.fixture
def run(book, tmp_path, capsys):
    write_xlsx(tmp_path / "cells.xlsx", sheet_xml(CELL_ROWS))

    def _run(*args):
        cwd = os.getcwd()
        os.chdir(tmp_path)
        try:
            code = main(["--color", "never", "--stats", *args])
        finally:
            os.chdir(cwd)
        out, err = capsys.readouterr()
        return code, out, err

    return _run


def test_table(run):
    code, out, _ = run()
    assert code == 0
    lines = out.splitlines()
    assert lines[0].split() == ["FILE", "SIZE", "SHEETS", "CELLS", "FORMULAS", "UNIQUE", "VOLATILE", "ARRAY",
                                "ERRORS", "VBA", "EXT"]
    # Sizes depend on the zip encoder; compare everything else ("8.9 KB" is two tokens).
    rows = [line.split() for line in lines[1:]]
    assert [r[:1] + r[3:] for r in rows] == [
        ["book.xlsx", "2/1", "9", "2", "2", "0", "0", "1", "yes", "1"],
        ["cells.xlsx", "1", "13", "10", "7", "2", "1", "2", "no", "0"],
        ["TOTAL", "3", "22", "12", "9", "2", "1", "3", "1", "1"],
    ]


def test_by_sheet_and_scope(run):
    _, out, _ = run("--by", "sheet", "--no-hidden", "book.xlsx")
    assert out.splitlines() == [
        "book.xlsx",
        "  SHEET  RANGE  CELLS  FORMULAS  UNIQUE  VOLATILE  ARRAY  ERRORS  CF  DV",
        "  Main   A1:D3      8         2       2         0      0       0   2   1",
    ]


def test_json_csv_cards(run):
    _, out, _ = run("--json", "book.xlsx")
    rec = json.loads(out)
    assert (rec["sheets"], rec["vba"], rec["charts"], rec["custom_functions"], rec["lambdas"]) == (
        2, True, 1, ["GetRate"], ["MyLambda"])
    _, out, _ = run("--csv", "--by", "sheet", "cells.xlsx")
    header, row = out.splitlines()
    assert header.startswith("file,sheet,hidden,range,cells,values,formulas,unique_formulas")
    assert "INDIRECT 1; NOW 1; RANDARRAY 1" in row or "NOW 1; INDIRECT 1; RANDARRAY 1" in row
    _, out, _ = run("-p", "cells.xlsx")
    assert "  volatile   2 formulas: " in out
    assert "  errors     2 cells: #REF! 1, #DIV/0! 1" in out


def test_bad_combinations(run):
    for args in (["-f", "SUM"], ["-C", "1"], ["--sort", "name"], ["--list-funcs"]):
        with pytest.raises(SystemExit):
            run(*args)
