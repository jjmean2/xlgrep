"""Formulas and text outside cells: defined names, conditional formats, validations, notes."""

import json
import os
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from xlgrep.cli import main
from xlgrep.objects import read_objects

# Excel 2010+ stores rules that reference other sheets in x14 extension blocks,
# which openpyxl can't write; patched in after saving.
X14_EXT = (
    '<extLst><ext uri="{78C0D931-6437-407d-A8EE-F0AAD7539E65}" '
    'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
    "<x14:conditionalFormattings><x14:conditionalFormatting "
    'xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
    '<x14:cfRule type="expression" priority="9" id="{00000000-0000-0000-0000-000000000001}">'
    "<xm:f>_xlfn.XLOOKUP($A1,Codes!A:A,Codes!B:B)=1</xm:f></x14:cfRule>"
    "<xm:sqref>D1:D20</xm:sqref></x14:conditionalFormatting></x14:conditionalFormattings>"
    "</ext>"
    '<ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
    'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
    '<x14:dataValidations count="1" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
    '<x14:dataValidation type="list"><x14:formula1><xm:f>Codes!$A$1:$A$9</xm:f></x14:formula1>'
    "<xm:sqref>E1:E5</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>"
)


@pytest.fixture
def objects_book(tmp_path: Path) -> Path:
    path = tmp_path / "rules.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "=SUM(B1:B3)"
    ws["C3"] = 1
    ws["C3"].comment = Comment("check the VLOOKUP\nbelow", "me")
    codes = wb.create_sheet("Codes")
    codes["A1"] = "x"
    hidden = wb.create_sheet("Secret")
    hidden.sheet_state = "hidden"
    hidden.conditional_formatting.add("A1:A5", FormulaRule(formula=["VLOOKUP(A1,Codes!A:B,2,0)"]))

    wb.defined_names["TaxRate"] = DefinedName("TaxRate", attr_text="Summary!$B$1")
    wb.defined_names["Lookup"] = DefinedName("Lookup", attr_text="VLOOKUP(Summary!A1,Codes!A:B,2,FALSE)")
    wb.defined_names["_xlfn.XLOOKUP"] = DefinedName("_xlfn.XLOOKUP", attr_text="#NAME?", hidden=True)
    ws.defined_names["Local"] = DefinedName("Local", attr_text="SUM(Summary!A1:A3)")
    ws.print_area = "A1:C3"  # becomes _xlnm.Print_Area, should be skipped

    ws.conditional_formatting.add("B2:B50 F2:F50", FormulaRule(formula=["VLOOKUP($A2,Codes!A:B,2,0)>100"]))
    ws.conditional_formatting.add("C1:C9", CellIsRule(operator="between", formula=["1", "SUM($A$1:$A$3)"]))
    dv = DataValidation(type="custom", formula1="COUNTIF($A:$A,A1)=1")
    dv.add("A1:A100")
    ws.add_data_validation(dv)
    wb.save(path)

    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"</worksheet>", X14_EXT.encode() + b"</worksheet>")
            dst.writestr(item, data)
    shutil.move(tmp, path)
    return path


@pytest.fixture
def run(objects_book, capsys):
    def _run(*args):
        cwd = os.getcwd()
        os.chdir(objects_book.parent)
        try:
            code = main(["--color", "never", *args])
        finally:
            os.chdir(cwd)
        out, err = capsys.readouterr()
        return code, out, err

    return _run


def test_read_objects(objects_book):
    objs = read_objects(objects_book, {"name", "cf", "dv", "note"})
    assert [(s.name, s.hidden) for s in objs.sheets] == [("Summary", False), ("Codes", False), ("Secret", True)]
    assert [(o.ref, o.text) for o in objs.workbook_names] == [
        ("TaxRate", "=Summary!$B$1"),
        ("Lookup", "=VLOOKUP(Summary!A1,Codes!A:B,2,FALSE)"),
    ]
    summary = [(o.object, o.location, o.text) for o in objs.by_sheet["Summary"]]
    assert summary == [
        ("cf", "Summary!B2:B50,F2:F50", "=VLOOKUP($A2,Codes!A:B,2,0)>100"),
        ("cf", "Summary!C1:C9", "=1"),
        ("cf", "Summary!C1:C9", "=SUM($A$1:$A$3)"),
        ("dv", "Summary!A1:A100", "=COUNTIF($A:$A,A1)=1"),
        ("cf", "Summary!D1:D20", "=XLOOKUP($A1,Codes!A:A,Codes!B:B)=1"),
        ("dv", "Summary!E1:E5", "=Codes!$A$1:$A$9"),
        ("name", "Summary!Local", "=SUM(Summary!A1:A3)"),
        ("note", "Summary!C3", "check the VLOOKUP\nbelow"),
    ]


def test_func_search_finds_objects(run):
    code, out, _ = run("-f", "VLOOKUP")
    assert code == 0
    assert out.splitlines() == [
        "rules.xlsx:Summary!B2:B50,F2:F50#cf:=VLOOKUP($A2,Codes!A:B,2,0)>100",
        "rules.xlsx:Secret!A1:A5#cf:=VLOOKUP(A1,Codes!A:B,2,0)",
        "rules.xlsx:Lookup#name:=VLOOKUP(Summary!A1,Codes!A:B,2,FALSE)",
    ]


def test_value_search_finds_notes_only(run):
    _, out, _ = run("--in", "value", "VLOOKUP")
    assert out == "rules.xlsx:Summary!C3#note:check the VLOOKUP\\nbelow\n"


def test_objects_option(run):
    _, out, _ = run("--objects", "cells", "SUM")
    assert out == "rules.xlsx:Summary!A1:=SUM(B1:B3)\n"
    _, out, _ = run("--objects", "dv,names", "-c", ".")
    assert out == "rules.xlsx:5\n"  # 2 dv + TaxRate, Lookup, Local


def test_sheet_and_range_filters(run):
    _, out, _ = run("--sheet", "Summary", "--objects", "names", ".")
    assert out == "rules.xlsx:Summary!Local#name:=SUM(Summary!A1:A3)\n"
    _, out, _ = run("--range", "F10", "--objects", "cf,dv,names", ".")
    assert out == "rules.xlsx:Summary!B2:B50,F2:F50#cf:=VLOOKUP($A2,Codes!A:B,2,0)>100\n"
    _, out, _ = run("--no-hidden", "-c", "-f", "VLOOKUP")
    assert out == "rules.xlsx:2\n"


def test_json_and_csv(run):
    _, out, _ = run("--json", "--objects", "dv", "COUNTIF")
    assert json.loads(out) == {
        "file": "rules.xlsx",
        "object": "dv",
        "sheet": "Summary",
        "ref": "A1:A100",
        "kind": "formula",
        "text": "=COUNTIF($A:$A,A1)=1",
        "matches": [{"start": 1, "end": 8, "text": "COUNTIF"}],
        "dv_type": "custom",
        "part": "formula1",
    }
    _, out, _ = run("--csv", "-f", "XLOOKUP")
    assert out.splitlines()[1] == 'rules.xlsx,Summary,cf,D1:D20,formula,"=XLOOKUP($A1,Codes!A:A,Codes!B:B)=1",'


def test_pretty_lists_objects_under_sheet(run):
    _, out, _ = run("-p", "SUM")
    assert out.splitlines() == [
        "rules.xlsx",
        "  Summary",
        "      │ A           │ B",
        "  ▶ 1 │ =SUM(B1:B3) │",
        "",
        "    #cf    C1:C9  =SUM($A$1:$A$3)",
        "    #name  Local  =SUM(Summary!A1:A3)",
        "  3 matches",
    ]
