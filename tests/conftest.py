from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest


def make_workbook(
    path: Path,
    sheets: dict[str, dict[str, object]],
    cached: dict[tuple[str, str], object] | None = None,
    hidden: tuple[str, ...] = (),
) -> Path:
    """Create a workbook. ``cached`` sets the stored result of formula cells.

    openpyxl never writes formula results, so they're patched into the sheet XML
    afterwards, the way Excel would have saved them.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cells in sheets.items():
        ws = wb.create_sheet(name)
        for addr, value in cells.items():
            ws[addr] = value
        if name in hidden:
            ws.sheet_state = "hidden"
    wb.save(path)

    if cached:
        names = list(sheets)
        tmp = path.with_suffix(".tmp")
        with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                data = src.read(item.filename)
                m = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
                if m:
                    sheet = names[int(m.group(1)) - 1]
                    xml = data.decode()
                    for (s, addr), value in cached.items():
                        if s == sheet:
                            xml = _patch_cached(xml, addr, value)
                    data = xml.encode()
                dst.writestr(item, data)
        shutil.move(tmp, path)
    return path


def _patch_cached(xml: str, addr: str, value: object) -> str:
    pattern = re.compile(rf'<c r="{addr}"([^>]*)>(<f[^>]*>.*?</f>)<v ?/>(</c>)')
    if isinstance(value, str):
        repl = rf'<c r="{addr}"\1 t="str">\2<v>{value}</v>\3'
    elif isinstance(value, bool):
        repl = rf'<c r="{addr}"\1 t="b">\2<v>{int(value)}</v>\3'
    else:
        repl = rf'<c r="{addr}"\1>\2<v>{value}</v>\3'
    new, n = pattern.subn(repl, xml)
    assert n == 1, f"no formula cell {addr} to patch"
    return new


@pytest.fixture
def sample_dir(tmp_path: Path) -> Path:
    """A small tree of workbooks used across CLI tests."""
    root = tmp_path / "data"
    (root / "sub").mkdir(parents=True)
    make_workbook(
        root / "sales.xlsx",
        {
            "Summary": {
                "A1": "지역", "B1": "매출액", "C1": "비고",
                "A2": "서울", "B2": "=VLOOKUP(A2,Data!A:D,4,FALSE)", "C2": "ok",
                "A3": "부산", "B3": "=VLOOKUP(A3,Data!A:D,4,FALSE)",
                "A4": "합계", "B4": "=SUM(B2:B3)",
                "B6": '="VLOOKUP is not called here"',
            },
            "Raw Data": {
                "A1": "key", "B1": "=_xlfn.XLOOKUP(A1,C:C,D:D)",
                "A2": "line one\nline two",
            },
        },
        cached={("Summary", "B2"): 1200, ("Summary", "B3"): 980, ("Summary", "B4"): 2180,
                ("Summary", "B6"): "VLOOKUP is not called here", ("Raw Data", "B1"): "#N/A"},
    )
    make_workbook(
        root / "sub" / "budget.xlsx",
        {"Sheet1": {"A1": "rate", "B1": "=vlookup(A1,Rates,2)", "C1": 3.5}, "Secret": {"A1": "VLOOKUP note"}},
        hidden=("Secret",),
    )
    (root / "~$sales.xlsx").write_bytes(b"lock file")
    (root / "notes.txt").write_text("VLOOKUP")
    return root
