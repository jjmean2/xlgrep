"""--list-funcs: function usage summary."""

import json
import os
from pathlib import Path

import openpyxl
import pytest
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName

from xlgrep.cli import main
from xlgrep.funcs import BUILTIN, CUSTOM, LAMBDA, classify, function_calls


def test_function_calls_skips_strings_and_keeps_prefixes():
    assert function_calls('=IFERROR(VLOOKUP(A1,"SUM(x)",2,0),_xlfn.XLOOKUP(B1,C:C,D:D))') == [
        "IFERROR", "VLOOKUP", "_xlfn.XLOOKUP",
    ]
    assert function_calls("=SUM(SUM(A1),B1)") == ["SUM", "SUM"]
    assert function_calls("={1,2;3,4}") == []


def test_function_calls_falls_back_on_unbalanced_formula():
    assert function_calls('=SUM(A1,"x(")+MAX(B1') == ["SUM", "MAX"]


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("vlookup", ("VLOOKUP", BUILTIN)),
        ("_xlfn.XLOOKUP", ("XLOOKUP", BUILTIN)),
        ("_xlfn._xlws.FILTER", ("FILTER", BUILTIN)),
        ("TABLE", ("TABLE", BUILTIN)),
        ("MyLambda", ("MyLambda", LAMBDA)),
        ("GetRate", ("GetRate", CUSTOM)),
        ("_xll.AddinFn", ("AddinFn", CUSTOM)),
    ],
)
def test_classify(raw, expected):
    assert classify(raw, {"MYLAMBDA"}) == expected


@pytest.fixture
def run(tmp_path: Path, capsys):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = '=IFERROR(VLOOKUP(B1,Data!A:B,2,0),"SUM(x)")'
    ws["A2"] = "=SUM(SUM(A1),_xlfn.XLOOKUP(B2,C:C,D:D))"
    ws["A3"] = "=MyLambda(3)+GetRate(A1)"
    wb.defined_names["MyLambda"] = DefinedName("MyLambda", attr_text="_xlfn.LAMBDA(_xlpm.x,_xlpm.x*2)")
    ws.conditional_formatting.add("B1:B9", FormulaRule(formula=["VLOOKUP($A1,Data!A:B,2,0)>1"]))
    data = wb.create_sheet("Data")
    data["C1"] = "=vlookup(A1,A:B,2,0)"
    wb.save(tmp_path / "a.xlsx")
    wb = openpyxl.Workbook()
    wb.active["A1"] = "=SUM(A2:A9)"
    wb.save(tmp_path / "b.xlsx")

    def _run(*args):
        cwd = os.getcwd()
        os.chdir(tmp_path)
        try:
            code = main(["--color", "never", "--list-funcs", *args])
        finally:
            os.chdir(cwd)
        out, err = capsys.readouterr()
        return code, out, err

    return _run


def test_table(run):
    code, out, _ = run()
    assert code == 0
    assert out.splitlines() == [
        "FUNCTION  CALLS  PLACES  FILES",
        "VLOOKUP       3       3      1",
        "SUM           3       2      2",
        "GetRate       1       1      1  custom",
        "IFERROR       1       1      1",
        "LAMBDA        1       1      1",
        "MyLambda      1       1      1  lambda",
        "XLOOKUP       1       1      1",
    ]


def test_by_file_with_func_filter(run):
    _, out, _ = run("--by", "file", "-f", "vlookup,sum")
    assert out.splitlines() == [
        "a.xlsx",
        "  FUNCTION  CALLS  PLACES",
        "  VLOOKUP       3       3",
        "  SUM           2       1",
        "",
        "b.xlsx",
        "  FUNCTION  CALLS  PLACES",
        "  SUM           1       1",
    ]


def test_by_sheet_with_pattern_and_sort(run):
    _, out, _ = run("--by", "sheet", "--sort", "name", "-e", "LOOK", "a.xlsx")
    assert out.splitlines() == [
        "a.xlsx",
        "  Summary",
        "    FUNCTION  CALLS  PLACES",
        "    VLOOKUP       2       2",
        "    XLOOKUP       1       1",
        "  Data",
        "    FUNCTION  CALLS  PLACES",
        "    VLOOKUP       1       1",
    ]


def test_scope_options_apply(run):
    _, out, _ = run("--objects", "cells", "-f", "VLOOKUP")
    assert out.splitlines()[1] == "VLOOKUP       2       2      1"
    _, out, _ = run("--sheet", "Data")
    assert out.splitlines()[1:] == ["VLOOKUP       1       1      1"]


def test_json_and_csv(run):
    _, out, _ = run("--json", "-e", "XL")
    assert [json.loads(line) for line in out.splitlines()] == [
        {"function": "XLOOKUP", "category": "builtin", "calls": 1, "places": 1, "files": 1},
    ]
    _, out, _ = run("--csv", "--by", "sheet", "-f", "MyLambda")
    assert out.splitlines() == [
        "file,sheet,function,category,calls,places",
        "a.xlsx,Summary,MyLambda,lambda,1,1",
    ]


def test_no_functions_and_bad_combinations(run):
    code, out, _ = run("-f", "NOPE")
    assert (code, out) == (1, "")
    with pytest.raises(SystemExit):
        run("-C", "1")
    with pytest.raises(SystemExit):
        main(["--by", "file", "SUM", "."])


def test_shape_cache_keeps_digit_function_names():
    from xlgrep.funcs import FuncCounter

    counter = FuncCounter(lambda name: True)
    for formula in ["=LOG(A1)", "=LOG10(A2)", "=DEC2BIN(A3)", "=LOG(A4)", "=SUM(A5:B5)"]:
        counter.add(None, "f", formula, set())
    stats = counter.groups[None]
    assert {k: (s.calls, s.places) for k, s in stats.items()} == {
        "LOG": (2, 2), "LOG10": (1, 1), "DEC2BIN": (1, 1), "SUM": (1, 1),
    }
